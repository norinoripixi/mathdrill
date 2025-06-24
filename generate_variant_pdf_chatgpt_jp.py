# 類題生成＆PDF出力用 Streamlit アプリ（ChatGPT + 日本語対応PDF）
import streamlit as st
import random
import re
from docx import Document
from io import BytesIO
from fpdf import FPDF
import openai
import os

st.title("📄 ChatGPTで類題プリント自動生成アプリ（日本語PDF対応）")

# APIキーの読み込み（安全対策付き）
if "OPENAI_API_KEY" in st.secrets:
    openai.api_key = st.secrets["OPENAI_API_KEY"]
else:
    st.error("OpenAI APIキーが設定されていません。secrets.toml に OPENAI_API_KEY を追加してください。")
    st.stop()

uploaded_file = st.file_uploader("Wordファイル（.docx）をアップロード", type="docx")

# ChatGPTを使って類題生成
def generate_variant_with_chatgpt(original):
    prompt = f"次の小学生向け算数の問題と同じ形式で、数字を変えた日本語の類題を1問作ってください：{original}"
    try:
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "あなたは日本の小学生向けの計算問題を作る教師です。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.7,
            max_tokens=100
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        return f"（変換失敗：{original}）"

# Word -> テキスト構造抽出 & ChatGPTで類題化
def extract_variant_paragraphs_with_chatgpt(docx_bytes):
    doc = Document(docx_bytes)
    variant_paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text == "":
            variant_paragraphs.append("")
        elif re.search(r'[0-9]', text) and any(op in text for op in ['＋', '+', '-', '×', '*', '÷', '/']):
            new_text = generate_variant_with_chatgpt(text)
            variant_paragraphs.append(new_text)
        else:
            variant_paragraphs.append(text)
    return variant_paragraphs

# PDF化（日本語フォント対応）
class PDF(FPDF):
    def header(self):
        self.set_font("IPAexGothic", size=14)
        self.cell(0, 10, "小４前期計算トレーニング（ChatGPT類題）", ln=True, align="C")
        self.ln(5)

    def add_lines(self, lines):
        self.set_font("IPAexGothic", size=12)
        for line in lines:
            self.multi_cell(0, 8, line)

if uploaded_file:
    st.success("ファイルを読み込みました。ChatGPTで類題を生成しています……")
    variant_lines = extract_variant_paragraphs_with_chatgpt(uploaded_file)

    pdf = PDF()
    pdf.add_font("IPAexGothic", "", "ipaexg.ttf", uni=True)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_lines(variant_lines)

    pdf_output = BytesIO()
    pdf.output(pdf_output)
    pdf_output.seek(0)

    st.download_button(
        label="📥 類題PDFをダウンロード",
        data=pdf_output,
        file_name="ChatGPT_類題トレーニング.pdf",
        mime="application/pdf"
    )
