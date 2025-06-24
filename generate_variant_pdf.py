# 類題生成＆PDF出力用 Streamlit アプリ
import streamlit as st
import random
import re
from docx import Document
from io import BytesIO
from fpdf import FPDF

st.title("📄 類題プリント自動生成アプリ")

uploaded_file = st.file_uploader("Wordファイル（.docx）をアップロード", type="docx")

# 数値の置換処理
def replace_numbers(text):
    def repl(m):
        original = int(m.group())
        if original < 10:
            return str(random.randint(2, 9))
        elif original < 100:
            return str(random.randint(10, 99))
        elif original < 1000:
            return str(random.randint(100, 999))
        else:
            return str(random.randint(1000, 9999))
    return re.sub(r'\d{1,5}', repl, text)

# Word -> テキスト構造抽出 & 数値置換
def extract_variant_paragraphs(docx_bytes):
    doc = Document(docx_bytes)
    variant_paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text == "":
            variant_paragraphs.append("")
        else:
            new_text = replace_numbers(text)
            variant_paragraphs.append(new_text)
    return variant_paragraphs

# PDF化
class PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 14)
        self.cell(0, 10, "小４前期計算トレーニング（類題）", ln=True, align="C")
        self.ln(5)

    def add_lines(self, lines):
        self.set_font("Arial", size=12)
        for line in lines:
            self.multi_cell(0, 8, line)

if uploaded_file:
    st.success("ファイルを読み込みました。類題を生成しています……")
    variant_lines = extract_variant_paragraphs(uploaded_file)

    pdf = PDF()
    pdf.add_page()
    pdf.add_lines(variant_lines)

    pdf_output = BytesIO()
    pdf.output(pdf_output)
    pdf_output.seek(0)

    st.download_button(
        label="📥 類題PDFをダウンロード",
        data=pdf_output,
        file_name="類題トレーニング.pdf",
        mime="application/pdf"
    )
